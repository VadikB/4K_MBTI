"""Build the SLS06 methodology review package from immutable repository inputs."""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import subprocess
import zipfile
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SALES = ROOT / "assessment_definitions/cases/competencies_4k/sales-pilot/0.1"
M6_BLIND = ROOT / "tests/fixtures/m6_gc/v1/blind_review"
WORKBOOK_BUILDER = ROOT / "scripts/build_sls06_workbooks.mjs"


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def set_cell_shading(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    props.append(shading)


def set_cell_margins(cell, value: int = 100) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge in ("top", "start", "bottom", "end"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(22)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor(0, 0, 0)


def add_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.8)
    table.columns[1].width = Inches(5.9)
    for index, heading in enumerate(("Поле", "Фактическое содержание")):
        cell = table.rows[0].cells[index]
        cell.text = heading
        set_cell_shading(cell, "1F4E78")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.bold = True
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value or "Не задано в источнике"
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    doc.add_paragraph()


def collect() -> dict:
    package = json.loads((SALES / "case-package.json").read_text())
    mapping = json.loads((SALES / "mapping-report.json").read_text())
    rules = json.loads((SALES / "execution-rules.json").read_text())
    previews = json.loads((SALES / "participant-preview.json").read_text())
    rule_by_case = rules["cases"]
    preview_by_case = {item["case_id"]: item for item in previews}
    map_by_case = {item["case_id"]: item for item in mapping}
    cases = []
    for item in package["cases"]:
        case_id = item["case_id"]
        author_as = item["author_assessment_situation"]
        participant = item["participant_presentation"]
        mapped = map_by_case[case_id]
        indicators = []
        for target in item["indicator_targets"]:
            source = target.get("source", "")
            if isinstance(source, dict):
                source = source.get("file", "") + (source.get("pointer", "") or "")
            indicators.append({
                "indicator_id": target["indicator_id"],
                "skill_id": target.get("skill_id", ""),
                "component_id": target.get("component_id", ""),
                "name": target.get("name", target.get("observable_action", "")),
                "source": source,
                "version": target.get("version", target.get("m2_version", "")),
            })
        cases.append({
            "case_id": case_id,
            "case_version": item["version"],
            "case_status": item["status"],
            "title": item["title"],
            "as_id": author_as["AssessmentSituationID"],
            "as_version": mapped["as_version"],
            "as_status": author_as.get("Статус", author_as.get("Status", "TEST_DRAFT")),
            "base_role": item["base_role"],
            "active_counterpart": author_as.get("Действующий собеседник", author_as.get("Role", "")),
            "channel": author_as.get("Канал", author_as.get("Контекстный снимок", {}).get("Канал", "")),
            "time_minutes": author_as.get("Плановое время минут", author_as.get("Плановое время", {})),
            "trigger": participant.get("trigger", ""),
            "introduction": participant.get("introduction", ""),
            "initial_situation": participant.get("initial_situation", ""),
            "assessee_task": participant.get("assessee_task", ""),
            "conditions": participant.get("conditions_and_constraints", []),
            "participants": participant.get("participants_and_positions", []),
            "initial_materials": participant.get("initial_materials", []),
            "closed_materials": author_as.get("Скрытые сведения", author_as.get("Фиксированные сведения", [])),
            "scenario": author_as.get("Сценарий", []),
            "conditional_reactions": author_as.get("Условные реакции", author_as.get("Правила персонажа", [])),
            "completion": author_as.get("Завершение", author_as.get("Фиксация исполнения", {})),
            "development_boundaries": author_as.get("Границы развития", author_as.get("Допуск", "")),
            "m6_handoff": author_as.get("Передача в M6", author_as.get("Фиксация исполнения", {})),
            "indicators": indicators,
            "case_source": item["author_case_source"],
            "as_source": item["author_as_source"],
            "mapping_unresolved": mapped.get("unresolved", []),
            "runtime_unresolved": item.get("unresolved_decisions", []),
            "execution_rules": rule_by_case.get(case_id, {}),
            "participant_preview": preview_by_case.get(case_id, {}),
            "review_status": "NOT_REVIEWED",
            "methodologist_decision": "",
            "expert_decision": "",
        })
    return {
        "generated_on": date.today().isoformat(),
        "release_sha": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "test_url": "https://test.4k-assistant.ru/",
        "app_version": "2.0.6",
        "package": {
            "version": package["version"], "status": package["status"],
            "runtime_enabled": package["runtime_enabled"], "blockers": package["blockers"],
        },
        "mechanisms": {
            "m6_source": "M6 Evidence Assessment v1.8 FROZEN",
            "evidence_prompt": "m6_evidence/1.0.0 draft",
            "indicator_prompt": "m6_indicator_assessment/1.0.0 draft",
            "admission_prompt": "m6_substantive_admission/2.0.0 draft",
            "model": "NOT_RUN",
        },
        "cases": cases,
    }


def write_start(data: dict, out: Path) -> None:
    body = f"""SLS06 — ПАКЕТ СОДЕРЖАТЕЛЬНОЙ ПРИЁМКИ SALES-КЕЙСОВ

Release SHA: {data['release_sha']}
Дата сборки: {data['generated_on']}
Test URL: {data['test_url']}
Версия test при подготовке: {data['app_version']}

СТАТУС
Пакет технически подготовлен для review. Методологическая приёмка и независимое
экспертное заключение не выполнены. Все девять Case/AS имеют статус NOT_REVIEWED;
WORKING, TEST_DRAFT, build PASS и завершение чата не означают ADMITTED.

ПОРЯДОК
1. Прочитайте Cases_Review.docx: участническая часть отделена от закрытой служебной.
2. Сверьте Source_to_Runtime.xlsx и точные SHA-256 исходников.
3. Выполните сценарии из Test_Scenarios.txt на test только с назначенным доступом.
4. Запишите наблюдения и решения в Review_Protocol.xlsx.
5. Независимые эксперты работают с отдельными файлами gc_review_packets/ без
   ответов тестируемой модели и без заключения другого эксперта.

ДОСТУП
Назначенный доступ выдаёт владелец организации штатным способом. Пароли, токены и
персональные данные в архив не включены. Если доступ не назначен, используйте
offline-материал и отметьте authenticated run как NOT_RUN.

ВЫПОЛНЕНО
- статическая сверка девяти Case/AS, версий, источников и M2 IndicatorID;
- технические SLS01–05 и CI выпуска 2.0.6;
- подготовка blind-review форм и машинных checksum.

НЕ ВЫПОЛНЕНО
- человеческое заключение методолога;
- независимые экспертные GC и три обязательных повтора;
- real provider run для этого review-пакета;
- authenticated визуальный прогон девяти кейсов в рамках SLS06.

ЗАМЕЧАНИЯ
Записывайте их в Review_Protocol.xlsx. Не публикуйте закрытые карточки участникам и
не переносите содержимое архива в открытый URL.
"""
    (out / "START_REVIEW.txt").write_text(body)


def write_scenarios(data: dict, out: Path) -> None:
    lines = ["SLS06 — РУЧНЫЕ СЦЕНАРИИ ПРОВЕРКИ", "", "Общие правила:",
             "- фиксировать фактические реплики, раскрытия, остановки, возвраты и checksum trace;",
             "- не считать HTTP 200 визуальной проверкой;",
             "- не использовать рабочие данные или реальные персональные сведения;", ""]
    for case in data["cases"]:
        lines += [f"{case['case_id']} {case['case_version']} / {case['as_id']} {case['as_version']}",
                  f"Старт: {case['trigger']}",
                  "Проверить: стартовую проекцию; условия раскрытия; конечные ветви; завершение; "
                  "паузу/возврат; оценочное уточнение; отсутствие раскрытия закрытой карточки.",
                  "Ожидаемая запись: runs/<case_id>/ с transcript.txt и trace.json.",
                  "Фактический SLS06 run: NOT_RUN — требуется назначенный authenticated доступ.", ""]
    (out / "Test_Scenarios.txt").write_text("\n".join(lines))


def write_docx(data: dict, out: Path) -> None:
    # Source collection and unit validation must not require the optional
    # document-rendering dependency used only to assemble the review ZIP.
    global Document, Inches, OxmlElement, Pt, qn, RGBColor
    global WD_ALIGN_PARAGRAPH, WD_CELL_VERTICAL_ALIGNMENT, WD_SECTION
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    style_doc(doc)
    doc.add_heading("Проверка Sales кейсов", 0)
    intro = doc.add_paragraph()
    intro.add_run("Назначение. ").bold = True
    intro.add_run("Документ помогает методологу проверить перенос девяти Case/AS в runtime-пакет. "
                  "Он фиксирует факты реализации и не содержит человеческого решения о допуске.")
    add_table(doc, [("Release SHA", data["release_sha"]), ("Статус", "Технически готово к review; NOT_REVIEWED"),
                    ("Пакет", f"sales-pilot/{data['package']['version']}; runtime_enabled={data['package']['runtime_enabled']}"),
                    ("Оценочный контур", data["mechanisms"]["m6_source"]),
                    ("Real provider", "NOT_RUN")])
    for index, case in enumerate(data["cases"], 1):
        if index > 1:
            doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_heading(f"{case['case_id']}  {case['title']}", level=1)
        doc.add_heading("Участническая часть", level=2)
        add_table(doc, [
            ("Версии", f"Case {case['case_version']} {case['case_status']}; AS {case['as_id']} {case['as_version']} {case['as_status']}"),
            ("Роль", case["base_role"]), ("Trigger", text(case["trigger"])),
            ("Вводная", text(case["introduction"])), ("Ситуация", text(case["initial_situation"])),
            ("Задача", text(case["assessee_task"])), ("Условия", text(case["conditions"])),
            ("Участники", text(case["participants"])), ("Стартовые материалы", text(case["initial_materials"])),
        ])
        doc.add_heading("Служебная часть", level=2)
        add_table(doc, [
            ("Собеседник", text(case["active_counterpart"])), ("Канал", text(case["channel"])),
            ("Плановое время", text(case["time_minutes"])), ("Закрытые сведения", text(case["closed_materials"])),
            ("Условия раскрытия", text(case["conditional_reactions"])), ("Ветви", text(case["scenario"])),
            ("Завершение", text(case["completion"])), ("Границы", text(case["development_boundaries"])),
            ("Передача M6", text(case["m6_handoff"])),
        ])
        doc.add_heading("Точные цели M2", level=2)
        rows = [(x["indicator_id"], f"{x['name']} | {x['component_id']} | {x['skill_id']} | {x['version']}") for x in case["indicators"]]
        add_table(doc, rows)
        doc.add_heading("Статус проверки", level=2)
        add_table(doc, [("Технический статус", case["review_status"]),
                        ("Неоднозначности", text(case["mapping_unresolved"] + case["runtime_unresolved"])),
                        ("Решение методолога", "Не заполнено"), ("Решение эксперта", "Не заполнено")])
    doc.save(out / "Cases_Review.docx")


def write_csv_inputs(data: dict, out: Path) -> None:
    source_rows = []
    protocol_rows = []
    for case in data["cases"]:
        for indicator in case["indicators"]:
            source_rows.append([case["case_id"], case["case_version"], case["as_id"], case["as_version"],
                                indicator["indicator_id"], indicator["component_id"], indicator["skill_id"],
                                indicator["name"], indicator["source"], indicator["version"],
                                case["case_source"]["sha256"], case["as_source"]["sha256"],
                                "runtime_disabled", text(case["mapping_unresolved"] + case["runtime_unresolved"]),
                                "PENDING_HUMAN_REVIEW"])
        protocol_rows.append([case["case_id"], case["case_version"], case["as_id"], case["as_version"],
                              "Case/AS content and runtime presentation", "", "", "source Case/AS + M2/M6",
                              "", "", "", "", "", "NOT_REVIEWED"])
    with (out / "source_to_runtime.csv").open("w", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["CaseID", "CaseVersion", "ASID", "ASVersion", "IndicatorID", "ComponentID", "SkillID",
                         "IndicatorName", "M2Source", "M2Version", "CaseSHA256", "ASSHA256", "RuntimeState",
                         "Ambiguity", "AdmissionState"])
        writer.writerows(source_rows)
    with (out / "review_protocol.csv").open("w", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["CaseID", "CaseVersion", "ASID", "ASVersion", "ReviewItem", "ObservedDeviation",
                         "ExpectedBehavior", "Source", "Severity", "DecisionOrScope", "Author", "DecisionDate",
                         "RunReference", "Status"])
        writer.writerows(protocol_rows)


def copy_blind_packets(out: Path) -> None:
    destination = out / "gc_review_packets" / "m6_candidates_v1"
    shutil.copytree(M6_BLIND, destination)


def write_runs(data: dict, out: Path) -> None:
    runs = out / "runs"
    runs.mkdir()
    (runs / "README.txt").write_text(
        "Authenticated SLS06 case runs: NOT_RUN. Real provider runs: NOT_RUN.\n"
        "No transcript is fabricated. Add transcript.txt and trace.json with matching IDs/checksum after execution.\n"
    )
    (runs / "technical-release.json").write_text(json.dumps({
        "type": "technical_ci_release", "synthetic": False, "semantic_gc": False,
        "release_sha": data["release_sha"], "app_version": data["app_version"],
        "result": "PASS", "scope": "SLS01-SLS05 and S10-A technical CI/deploy; not methodology admission",
        "provider_model": "NOT_RUN",
    }, ensure_ascii=False, indent=2) + "\n")


def manifest_and_zip(data: dict, out: Path, zip_path: Path) -> None:
    files = []
    for path in sorted(p for p in out.rglob("*") if p.is_file() and p.name != "manifest.json"):
        files.append({"path": str(path.relative_to(out)), "sha256": sha256(path), "bytes": path.stat().st_size})
    manifest = {"kind": "sls06_methodology_review", "release_sha": data["release_sha"],
                "status": "READY_FOR_HUMAN_REVIEW", "human_admission": "NOT_REVIEWED", "files": files}
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in sorted(p for p in out.rglob("*") if p.is_file()):
            archive.write(path, path.relative_to(out.parent))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--node", required=True)
    args = parser.parse_args()
    data = collect()
    out = args.output_dir / f"Methodology_Review_{data['release_sha']}"
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    (out / "review-data.json").write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")
    write_start(data, out)
    write_scenarios(data, out)
    write_docx(data, out)
    write_csv_inputs(data, out)
    copy_blind_packets(out)
    write_runs(data, out)
    subprocess.run([args.node, str(WORKBOOK_BUILDER), str(out)], check=True)
    (out / "source_to_runtime.csv").unlink()
    (out / "review_protocol.csv").unlink()
    for diagnostic in out.glob("*.inspect.ndjson"):
        diagnostic.unlink()
    zip_path = args.output_dir / f"Methodology_Review_{data['release_sha']}.zip"
    manifest_and_zip(data, out, zip_path)
    print(zip_path)


if __name__ == "__main__":
    main()
